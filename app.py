import io
import os
import pickle
import re
import threading
import time
from datetime import datetime
from zoneinfo import ZoneInfo

import openpyxl
import pandas as pd
import plotly.graph_objects as go
import requests
import streamlit as st

# -----------------------------------------------------------------------------
# PAGE CONFIGURATION
# -----------------------------------------------------------------------------
st.set_page_config(
    page_title="Overdue Works Dashboard",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
        .block-container {
            max-width: 100% !important;
            padding-left: 1.2rem !important;
            padding-right: 1.2rem !important;
            padding-top: 2.8rem !important;
            padding-bottom: 2rem !important;
        }
        [data-testid="stPlotlyChart"] {
            width: 100% !important;
        }
    </style>
    """,
    unsafe_allow_html=True,
)

# -----------------------------------------------------------------------------
# CONSTANTS & TARGET SHEETS
# -----------------------------------------------------------------------------
TARGET_SHEETS = [
    "4702_MOT",
    "4702-New Tanks",
    "4702-LIS",
    "4702_A&P",
    "PMKSY-SMI-RRR",
    "4702_Paschima Vahini",
    "4711__Other Flood Control",
    "4702-Kharland",
    "NABARD - New Tank",
    "SDMF",
    "KMERC",
    "NABARD - A&P",
    "ATAL BHUJAL",
    "NABARD_MOT",
    "NABARD - LIS",
    "other dept funds",
    "SCP (LIS)",
    "SCP (Borewell)",
    "SCP (Checkdam)",
    "TSP (Checkdam)",
    "TSP (Borewell)",
    "TSP (LIS)",
    "NABARD - Other Flood Control",
]

COL_SO = "ಶಾಖಾಧಿಕಾರಿಯ ಹೆಸರು"
COL_AEE = "ಸಹಾಯಕ ಕಾರ್ಯಪಾಲಕ ಇಂಜಿನಿಯರ್ ಹೆಸರು"
COL_TARGET_DATE = "ಕಾಮಗಾರಿ ಮುಕ್ತಾಯಗೊಳಿಸಬೇಕಾದ ದಿನಾಂಕ"
COL_ACTUAL_DATE = "ಮುಕ್ತಾಯವಾದ ದಿನಾಂಕ"
COL_STATUS_STD = "CP/CF/P/TBS/TP/ETS/OTH"

DEFAULT_SHEET_URL = st.secrets.get(
    "SHEET_URL",
    "https://docs.google.com/spreadsheets/d/13W04z-fdU3RhuVJ5vQFRpK4BBSQRJWfscPqpr7Oj2Y8/edit?gid=1055199869#gid=1055199869",
)

BACKUP_CACHE_PATH = "/tmp/mi_overdue_dashboard_backup.pkl"


# -----------------------------------------------------------------------------
# GLOBAL LOCK (Prevents simultaneous requests from colliding with 410 Gone)
# -----------------------------------------------------------------------------
@st.cache_resource
def get_download_lock() -> threading.Lock:
    return threading.Lock()


# -----------------------------------------------------------------------------
# HELPER FUNCTIONS
# -----------------------------------------------------------------------------
def normalize_text(val) -> str:
    """Collapse multiple spaces/newlines and strip text."""
    if val is None or pd.isna(val):
        return ""
    return re.sub(r"\s+", " ", str(val)).strip()


def format_cell_for_display(val) -> str:
    """Format raw Excel cell values cleanly for table display and export."""
    if val is None or pd.isna(val):
        return ""
    if isinstance(val, (datetime, pd.Timestamp)):
        return val.strftime("%d-%m-%Y")
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return normalize_text(val)


def extract_spreadsheet_key(sheet_url: str) -> str:
    """Extract the Google Spreadsheet ID from the URL."""
    match = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", sheet_url)
    if not match:
        raise ValueError("Invalid Google Sheet URL. Could not extract spreadsheet key.")
    return match.group(1)


def download_workbook_bytes(sheet_url: str) -> bytes:
    """
    Download the Google Sheet as .xlsx with retry logic, browser headers,
    180-second streaming timeout, and fallback export endpoints to avoid 410 Gone.
    """
    key = extract_spreadsheet_key(sheet_url)
    candidate_urls = [
        f"https://docs.google.com/spreadsheets/d/{key}/export?format=xlsx",
        f"https://docs.google.com/spreadsheets/d/{key}/export?format=xlsx&id={key}",
        f"https://drive.google.com/uc?export=download&id={key}",
    ]
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/128.0.0.0 Safari/537.36"
        ),
        "Accept": "*/*",
        "Cache-Control": "no-cache",
    }

    last_error = None
    for attempt in range(3):
        url = candidate_urls[attempt % len(candidate_urls)]
        # Add cache-buster on retries so Google generates a fresh export stream
        if attempt > 0:
            sep = "&" if "?" in url else "?"
            url = f"{url}{sep}t={int(time.time())}"

        try:
            with requests.Session() as session:
                with session.get(
                    url,
                    headers=headers,
                    timeout=(20, 180),
                    stream=True,
                    allow_redirects=True,
                ) as resp:
                    resp.raise_for_status()
                    chunks = []
                    for chunk in resp.iter_content(chunk_size=262144):
                        if chunk:
                            chunks.append(chunk)
                    content = b"".join(chunks)
                    if len(content) > 1000:
                        return content
        except Exception as err:
            last_error = err
            time.sleep(2 * (attempt + 1))

    raise RuntimeError(f"Failed after 3 download attempts. Last error: {last_error}")


def find_matching_column_indices(headers: list[str]) -> dict:
    """Find required 0-based column positions across sheets."""
    col_idx_map = {}
    for idx, col in enumerate(headers):
        norm_col = normalize_text(col)
        upper_col = norm_col.upper()
        compact_upper = re.sub(r"\s+", "", upper_col)

        if (COL_SO in norm_col or "ಶಾಖಾಧಿಕಾರಿ" in norm_col) and "so" not in col_idx_map:
            col_idx_map["so"] = idx
        elif (
            COL_AEE in norm_col or "ಸಹಾಯಕ ಕಾರ್ಯಪಾಲಕ ಇಂಜಿನಿಯರ್" in norm_col
        ) and "aee" not in col_idx_map:
            col_idx_map["aee"] = idx
        elif (
            COL_TARGET_DATE in norm_col or "ಮುಕ್ತಾಯಗೊಳಿಸಬೇಕಾದ ದಿನಾಂಕ" in norm_col
        ) and "target_date" not in col_idx_map:
            col_idx_map["target_date"] = idx
        elif (
            COL_ACTUAL_DATE in norm_col
            and "ಮುಕ್ತಾಯಗೊಳಿಸಬೇಕಾದ" not in norm_col
            and "actual_date" not in col_idx_map
        ):
            col_idx_map["actual_date"] = idx
        elif "CP/CF" in compact_upper and "status" not in col_idx_map:
            col_idx_map["status"] = idx

    return col_idx_map


def build_clean_unique_headers(raw_headers: list[str], col_idx_map: dict) -> list[str]:
    """Standardize key column names across sheets and deduplicate repeated headers."""
    standardized = list(raw_headers)
    if "so" in col_idx_map:
        standardized[col_idx_map["so"]] = COL_SO
    if "aee" in col_idx_map:
        standardized[col_idx_map["aee"]] = COL_AEE
    if "target_date" in col_idx_map:
        standardized[col_idx_map["target_date"]] = COL_TARGET_DATE
    if "actual_date" in col_idx_map:
        standardized[col_idx_map["actual_date"]] = COL_ACTUAL_DATE
    if "status" in col_idx_map:
        standardized[col_idx_map["status"]] = COL_STATUS_STD

    seen = {}
    unique_headers = []
    for idx, h in enumerate(standardized):
        base = h if h else f"Unnamed_Col_{idx + 1}"
        if base not in seen:
            seen[base] = 1
            unique_headers.append(base)
        else:
            seen[base] += 1
            unique_headers.append(f"{base} ({seen[base]})")
    return unique_headers


def parse_indian_date(series: pd.Series) -> pd.Series:
    """Safely parse dates (handling datetime objects and DD-MM-YYYY / DD.MM.YYYY strings)."""
    def _clean_date(val):
        if val is None or pd.isna(val):
            return pd.NaT
        if isinstance(val, (datetime, pd.Timestamp)):
            return pd.Timestamp(val.date())
        s = normalize_text(val)
        if not s or s.upper() in {"-", "--", "NIL", "NA", "N/A", "NONE", ""}:
            return pd.NaT
        s = re.sub(r"[./]", "-", s)
        return s

    cleaned = series.map(_clean_date)
    return pd.to_datetime(cleaned, dayfirst=True, errors="coerce")


def is_empty_date(series: pd.Series) -> pd.Series:
    """Check if the 'ಮುಕ್ತಾಯವಾದ ದಿನಾಂಕ' cell is empty or placeholder text."""
    def _check_empty(val):
        if val is None or pd.isna(val):
            return True
        s = normalize_text(val)
        return s.upper() in {"", "-", "--", "NIL", "NA", "N/A", "NONE", "NAT"}

    return series.map(_check_empty)


def is_valid_officer_name(name: str) -> bool:
    """Filter out empty cells, serial numbers, or summary rows like 'ಒಟ್ಟು' / 'Total'."""
    if not name or len(name) > 65:
        return False
    upper = name.upper()
    invalid_tokens = {
        "NAN", "NONE", "-", "--", "---", "NIL", "NA", "N/A", "0", "ಇಲ್ಲ", "ಖಾಲಿ"
    }
    if upper in invalid_tokens or name.isdigit():
        return False
    if "ಒಟ್ಟು" in name or "TOTAL" in upper or COL_SO in name or COL_AEE in name:
        return False
    return True


def shorten_label(name: str, max_len: int = 26) -> str:
    """Keep Y-axis labels compact so the chart always starts on the far left."""
    return name if len(name) <= max_len else name[: max_len - 1] + "…"


def convert_df_to_excel_bytes(df: pd.DataFrame) -> bytes:
    """Export a DataFrame to an in-memory .xlsx file."""
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Overdue_Works")
    return output.getvalue()


def parse_workbook_bytes(workbook_bytes: bytes):
    """Parse the 23 target worksheets using fast read-only OpenPyXL mode."""
    wb = openpyxl.load_workbook(
        io.BytesIO(workbook_bytes),
        read_only=True,
        data_only=True,
    )

    available_sheets = {
        normalize_text(name).lower(): name for name in wb.sheetnames
    }

    today_ist = pd.Timestamp(datetime.now(ZoneInfo("Asia/Kolkata")).date())

    compiled_rows = []
    full_overdue_rows = []
    sheet_diagnostics = []

    for target in TARGET_SHEETS:
        norm_target = normalize_text(target).lower()
        actual_sheet_name = available_sheets.get(norm_target)

        if not actual_sheet_name:
            sheet_diagnostics.append(
                {
                    "Sheet": target,
                    "Status": "Missing in Workbook",
                    "Rows Loaded": 0,
                    "Overdue Rows": 0,
                }
            )
            continue

        try:
            ws = wb[actual_sheet_name]
            rows_list = list(ws.iter_rows(values_only=True))
            if not rows_list:
                sheet_diagnostics.append(
                    {
                        "Sheet": target,
                        "Status": "Empty Sheet",
                        "Rows Loaded": 0,
                        "Overdue Rows": 0,
                    }
                )
                continue

            raw_df = pd.DataFrame(rows_list)

            header_idx = None
            for idx in range(min(25, len(raw_df))):
                row_text = " | ".join(normalize_text(v) for v in raw_df.iloc[idx].values)
                if (
                    COL_SO in row_text
                    or COL_AEE in row_text
                    or "ಶಾಖಾಧಿಕಾರಿ" in row_text
                ):
                    header_idx = idx
                    break

            if header_idx is None:
                sheet_diagnostics.append(
                    {
                        "Sheet": target,
                        "Status": "Header Row Not Found",
                        "Rows Loaded": 0,
                        "Overdue Rows": 0,
                    }
                )
                continue

            raw_headers = [normalize_text(col) for col in raw_df.iloc[header_idx].values]
            col_idx_map = find_matching_column_indices(raw_headers)

            required_keys = {"so", "aee", "target_date", "actual_date", "status"}
            missing_keys = required_keys - set(col_idx_map.keys())

            if missing_keys:
                sheet_diagnostics.append(
                    {
                        "Sheet": target,
                        "Status": f"Missing Columns: {', '.join(sorted(missing_keys))}",
                        "Rows Loaded": 0,
                        "Overdue Rows": 0,
                    }
                )
                continue

            df = raw_df.iloc[header_idx + 1 :].reset_index(drop=True).copy()

            s_so = df.iloc[:, col_idx_map["so"]].map(normalize_text)
            s_aee = df.iloc[:, col_idx_map["aee"]].map(normalize_text)
            s_status = df.iloc[:, col_idx_map["status"]].map(normalize_text)
            s_target_raw = df.iloc[:, col_idx_map["target_date"]]
            s_actual_raw = df.iloc[:, col_idx_map["actual_date"]]

            target_dt = parse_indian_date(s_target_raw)
            actual_empty = is_empty_date(s_actual_raw)
            is_status_p = s_status.str.upper() == "P"

            is_overdue = (
                is_status_p
                & target_dt.notna()
                & (target_dt < today_ist)
                & actual_empty
            )

            valid_row_mask = (s_status != "") | (s_so != "") | (s_aee != "")

            sub = pd.DataFrame(
                {
                    "sheet_name": target,
                    "so_name": s_so,
                    "aee_name": s_aee,
                    "status_raw": s_status,
                    "target_date_raw": s_target_raw.map(format_cell_for_display),
                    "actual_date_raw": s_actual_raw.map(format_cell_for_display),
                    "target_dt": target_dt,
                    "actual_empty": actual_empty,
                    "is_status_p": is_status_p,
                    "is_overdue": is_overdue,
                }
            )
            sub = sub[valid_row_mask].reset_index(drop=True)
            compiled_rows.append(sub)

            overdue_mask = valid_row_mask & is_overdue
            if overdue_mask.any():
                unique_headers = build_clean_unique_headers(raw_headers, col_idx_map)
                full_df = df.loc[overdue_mask].copy()
                full_df.columns = unique_headers

                unnamed_cols = [
                    c for c in full_df.columns if c.startswith("Unnamed_Col_")
                ]
                for u_col in unnamed_cols:
                    if full_df[u_col].map(format_cell_for_display).eq("").all():
                        full_df.drop(columns=[u_col], inplace=True)

                for col in full_df.columns:
                    full_df[col] = full_df[col].map(format_cell_for_display)

                full_df.insert(0, "Worksheet / ಯೋಜನೆ", target)
                full_df["_so_name"] = s_so.loc[overdue_mask].values
                full_df["_aee_name"] = s_aee.loc[overdue_mask].values

                full_overdue_rows.append(full_df.reset_index(drop=True))

            sheet_diagnostics.append(
                {
                    "Sheet": target,
                    "Status": "Loaded",
                    "Rows Loaded": len(sub),
                    "Overdue Rows": int(overdue_mask.sum()),
                }
            )

        except Exception as sheet_err:
            sheet_diagnostics.append(
                {
                    "Sheet": target,
                    "Status": f"Error: {sheet_err}",
                    "Rows Loaded": 0,
                    "Overdue Rows": 0,
                }
            )

    wb.close()

    master_df = (
        pd.concat(compiled_rows, ignore_index=True)
        if compiled_rows
        else pd.DataFrame()
    )
    overdue_full_df = (
        pd.concat(full_overdue_rows, ignore_index=True).fillna("")
        if full_overdue_rows
        else pd.DataFrame()
    )
    diag_df = pd.DataFrame(sheet_diagnostics)
    return master_df, overdue_full_df, diag_df, today_ist.strftime("%d-%m-%Y")


@st.cache_data(ttl=900, show_spinner=False)
def load_and_process_sheets_cached(sheet_url: str):
    """Download and parse with cache + disk backup."""
    workbook_bytes = download_workbook_bytes(sheet_url)
    result = parse_workbook_bytes(workbook_bytes)
    try:
        with open(BACKUP_CACHE_PATH, "wb") as f:
            pickle.dump(result, f)
    except Exception:
        pass
    return result


def get_dashboard_data(sheet_url: str, uploaded_file=None):
    """
    Thread-safe wrapper so multiple simultaneous users never trigger
    colliding Google Sheet downloads (410 Gone / Read Timeout).
    """
    if uploaded_file is not None:
        return parse_workbook_bytes(uploaded_file.getvalue()), False

    lock = get_download_lock()
    with lock:
        try:
            return load_and_process_sheets_cached(sheet_url), False
        except Exception as live_err:
            # If a backup exists from an earlier successful load, serve it seamlessly
            if os.path.exists(BACKUP_CACHE_PATH):
                try:
                    with open(BACKUP_CACHE_PATH, "rb") as f:
                        return pickle.load(f), True
                except Exception:
                    pass
            raise live_err


# -----------------------------------------------------------------------------
# SIDEBAR CONTROLS
# -----------------------------------------------------------------------------
st.sidebar.header("⚙️ Dashboard Controls")

sheet_url_input = st.sidebar.text_input(
    "Google Spreadsheet Link",
    value=DEFAULT_SHEET_URL,
    help="Ensure the Google Sheet is set to 'Anyone with the link can view'.",
)

uploaded_xlsx = st.sidebar.file_uploader(
    "Or Upload Workbook (.xlsx) Directly (Optional Backup)",
    type=["xlsx"],
    help="If Google Sheets is slow during a live presentation, upload the downloaded .xlsx file here.",
)

denominator_mode = st.sidebar.radio(
    "Total Works Assigned Count Basis:",
    options=[
        "All Works Assigned (All Statuses)",
        "Only Ongoing Works (Status = 'Progress')",
    ],
    index=0,
)

chart_style = st.sidebar.radio(
    "Bar Width Mode:",
    options=[
        "Overlay (Overdue inside Total Works)",
        "Scale by Overdue Works (Full-Width Bars)",
        "Scale by Overdue Percentage (%)",
    ],
    index=0,
    help="Choose how the horizontal bars are scaled across the page.",
)

hide_zero_overdue = st.sidebar.checkbox(
    "Show only officers with Overdue > 0",
    value=False,
)

if st.sidebar.button("🔄 Refresh Live Data"):
    st.cache_data.clear()

# -----------------------------------------------------------------------------
# MAIN UI
# -----------------------------------------------------------------------------
st.title("📊 Minor Irrigation — Overdue Works Dashboard")

if not sheet_url_input and uploaded_xlsx is None:
    st.info("Please paste your Google Sheet link or upload the .xlsx file in the sidebar.")
    st.stop()

with st.spinner("Fetching and processing 23 worksheets..."):
    try:
        (master_df, overdue_full_df, diag_df, today_str), used_backup = get_dashboard_data(
            sheet_url_input, uploaded_xlsx
        )
        if used_backup:
            st.warning(
                "⚡ Google Sheets took too long to respond to concurrent requests — displaying the most recent cached snapshot."
            )
    except Exception as e:
        st.error(
            f"Google Sheets is temporarily throttling live export requests.\n\n"
            f"**Quick Fix:** Wait 10 seconds and reload, or upload the `.xlsx` file in the left sidebar.\n\n"
            f"Technical Details: {e}"
        )
        st.stop()

if master_df.empty:
    st.warning("No matching data rows were found across the target worksheets.")
    st.dataframe(diag_df, use_container_width=True)
    st.stop()

# -----------------------------------------------------------------------------
# TOP CONTROLS: OFFICER VIEW & SORT ORDER
# -----------------------------------------------------------------------------
ctrl_col1, ctrl_col2 = st.columns(2)

with ctrl_col1:
    view_choice = st.radio(
        "ವೀಕ್ಷಣೆ ಆಯ್ಕೆಮಾಡಿ / Select Officer View:",
        options=[COL_SO, COL_AEE],
        horizontal=True,
    )

with ctrl_col2:
    sort_choice = st.radio(
        "ಕ್ರಮಬದ್ಧಗೊಳಿಸಿ / Sort Chart & Table By:",
        options=[
            "Highest Overdue Count at Top",
            "Highest Overdue Percentage (%) at Top",
        ],
        horizontal=True,
    )

group_col = "so_name" if view_choice == COL_SO else "aee_name"
full_filter_col = "_so_name" if view_choice == COL_SO else "_aee_name"

valid_df = master_df[master_df[group_col].map(is_valid_officer_name)].copy()

if denominator_mode.startswith("Only Ongoing Works"):
    valid_df = valid_df[valid_df["is_status_p"]].copy()

summary_df = (
    valid_df.groupby(group_col, as_index=False)
    .agg(
        total_works=("is_overdue", "count"),
        ongoing_p_works=("is_status_p", "sum"),
        overdue_works=("is_overdue", "sum"),
    )
    .rename(columns={group_col: view_choice})
)

summary_df["overdue_pct"] = (
    (summary_df["overdue_works"] / summary_df["total_works"]) * 100
).round(1).fillna(0.0)

if sort_choice == "Highest Overdue Percentage (%) at Top":
    summary_df = summary_df.sort_values(
        by=["overdue_pct", "overdue_works", "total_works"],
        ascending=[False, False, False],
    ).reset_index(drop=True)
else:
    summary_df = summary_df.sort_values(
        by=["overdue_works", "overdue_pct", "total_works"],
        ascending=[False, False, False],
    ).reset_index(drop=True)

# -----------------------------------------------------------------------------
# KPI SUMMARY METRICS
# -----------------------------------------------------------------------------
total_assigned_all = int(summary_df["total_works"].sum())
total_overdue_all = int(summary_df["overdue_works"].sum())
overall_overdue_pct = (
    round((total_overdue_all / total_assigned_all) * 100, 1)
    if total_assigned_all > 0
    else 0.0
)

k1, k2, k3, k4 = st.columns(4)
k1.metric("📅 Comparison Date (Today)", today_str)
k2.metric(f"👥 Total Officers ({view_choice})", len(summary_df))
k3.metric("📋 Total Works Assigned", f"{total_assigned_all:,}")
k4.metric(
    "⚠️ Total Overdue Works",
    f"{total_overdue_all:,} ({overall_overdue_pct}%)",
)

st.divider()

# -----------------------------------------------------------------------------
# HORIZONTAL BAR CHART (FULL WIDTH & LEFT ALIGNED)
# -----------------------------------------------------------------------------
st.subheader(f"📈 Overdue Works by {view_choice} ({sort_choice})")

chart_source_df = (
    summary_df[summary_df["overdue_works"] > 0].copy()
    if hide_zero_overdue
    else summary_df.copy()
)

if chart_source_df.empty:
    st.info("No officers match the current filter.")
else:
    plot_df = chart_source_df.iloc[::-1].copy()
    plot_df["y_short"] = plot_df[view_choice].map(lambda s: shorten_label(s, 28))

    if sort_choice == "Highest Overdue Percentage (%) at Top":
        plot_df["bar_label"] = plot_df.apply(
            lambda r: (
                f"  <b>{r['overdue_pct']:.1f}% Overdue</b> "
                f"({int(r['overdue_works'])} Overdue / {int(r['total_works'])} Total)"
            ),
            axis=1,
        )
    else:
        plot_df["bar_label"] = plot_df.apply(
            lambda r: (
                f"  <b>{int(r['overdue_works'])} Overdue</b> / "
                f"{int(r['total_works'])} Total ({r['overdue_pct']:.1f}%)"
            ),
            axis=1,
        )

    fig = go.Figure()

    if chart_style == "Scale by Overdue Works (Full-Width Bars)":
        fig.add_trace(
            go.Bar(
                y=plot_df[view_choice],
                x=plot_df["overdue_works"],
                name="Overdue Works",
                orientation="h",
                marker=dict(
                    color="#EF4444",
                    line=dict(color="#DC2626", width=1),
                ),
                text=plot_df["bar_label"],
                textposition="outside",
                textfont=dict(size=13),
                cliponaxis=False,
                customdata=plot_df[
                    [view_choice, "total_works", "ongoing_p_works", "overdue_pct"]
                ],
                hovertemplate=(
                    f"<b>{view_choice}:</b> %{{customdata[0]}}<br>"
                    "<b>Overdue Works:</b> %{x}<br>"
                    "<b>Total Works Assigned:</b> %{customdata[1]}<br>"
                    "<b>Ongoing ('Progress') Works:</b> %{customdata[2]}<br>"
                    "<b>Overdue Percentage:</b> %{customdata[3]:.1f}%<extra></extra>"
                ),
            )
        )
        max_x = max(int(plot_df["overdue_works"].max()), 1)
        x_range_max = max_x * 1.30
        x_axis_title = "Number of Overdue Works"

    elif chart_style == "Scale by Overdue Percentage (%)":
        fig.add_trace(
            go.Bar(
                y=plot_df[view_choice],
                x=plot_df["overdue_pct"],
                name="Overdue %",
                orientation="h",
                marker=dict(
                    color="#EF4444",
                    line=dict(color="#DC2626", width=1),
                ),
                text=plot_df["bar_label"],
                textposition="outside",
                textfont=dict(size=13),
                cliponaxis=False,
                customdata=plot_df[
                    [view_choice, "total_works", "ongoing_p_works", "overdue_works"]
                ],
                hovertemplate=(
                    f"<b>{view_choice}:</b> %{{customdata[0]}}<br>"
                    "<b>Overdue Percentage:</b> %{x:.1f}%<br>"
                    "<b>Overdue Works:</b> %{customdata[3]}<br>"
                    "<b>Total Works Assigned:</b> %{customdata[1]}<br>"
                    "<b>Ongoing ('Progress') Works:</b> %{customdata[2]}<extra></extra>"
                ),
            )
        )
        max_x = max(float(plot_df["overdue_pct"].max()), 1.0)
        x_range_max = max_x * 1.30
        x_axis_title = "Overdue Percentage (%)"

    else:
        fig.add_trace(
            go.Bar(
                y=plot_df[view_choice],
                x=plot_df["total_works"],
                name="Total Works Assigned",
                orientation="h",
                marker=dict(
                    color="rgba(148, 163, 184, 0.30)",
                    line=dict(color="#64748B", width=1),
                ),
                text=plot_df["bar_label"],
                textposition="outside",
                textfont=dict(size=13),
                cliponaxis=False,
                customdata=plot_df[[view_choice]],
                hovertemplate=(
                    f"<b>{view_choice}:</b> %{{customdata[0]}}<br>"
                    "<b>Total Works Assigned:</b> %{x}<extra></extra>"
                ),
            )
        )
        fig.add_trace(
            go.Bar(
                y=plot_df[view_choice],
                x=plot_df["overdue_works"],
                name="Overdue Works",
                orientation="h",
                marker=dict(color="#EF4444"),
                customdata=plot_df[
                    [view_choice, "total_works", "ongoing_p_works", "overdue_pct"]
                ],
                hovertemplate=(
                    f"<b>{view_choice}:</b> %{{customdata[0]}}<br>"
                    "<b>Overdue Works:</b> %{x}<br>"
                    "<b>Total Works Assigned:</b> %{customdata[1]}<br>"
                    "<b>Ongoing ('Progress') Works:</b> %{customdata[2]}<br>"
                    "<b>Overdue Percentage:</b> %{customdata[3]:.1f}%<extra></extra>"
                ),
            )
        )
        max_x = max(int(plot_df["total_works"].max()), 1)
        x_range_max = max_x * 1.32
        x_axis_title = "Number of Works"

    chart_height = max(500, len(plot_df) * 36 + 90)

    fig.update_layout(
        barmode="overlay",
        height=chart_height,
        bargap=0.22,
        xaxis=dict(
            title=x_axis_title,
            range=[0, x_range_max],
            gridcolor="rgba(200, 200, 200, 0.2)",
            zeroline=True,
        ),
        yaxis=dict(
            title="",
            automargin=False,
            tickmode="array",
            tickvals=plot_df[view_choice],
            ticktext=plot_df["y_short"],
            tickfont=dict(size=13),
        ),
        legend=dict(
            orientation="h",
            yanchor="bottom",
            y=1.01,
            xanchor="left",
            x=0,
        ),
        margin=dict(l=210, r=40, t=40, b=40),
    )

    st.plotly_chart(fig, use_container_width=True)

# -----------------------------------------------------------------------------
# DETAILED SUMMARY TABLE
# -----------------------------------------------------------------------------
st.subheader("📋 Officer-Wise Summary Table")

display_table = summary_df.rename(
    columns={
        "total_works": "Total Works Assigned",
        "ongoing_p_works": "Ongoing Works ('Progress')",
        "overdue_works": "Overdue Works",
        "overdue_pct": "Overdue %",
    }
)
display_table.index = display_table.index + 1

st.dataframe(
    display_table.style.format({"Overdue %": "{:.1f}%"}),
    use_container_width=True,
)

st.divider()

# -----------------------------------------------------------------------------
# BOTTOM SECTION: OFFICER DROPDOWN & COMBINED OVERDUE WORKS TABLE + DOWNLOAD
# -----------------------------------------------------------------------------
st.subheader(
    f"📂 ಅಧಿಕಾರಿಗಳ ವಾರು ಅವಧಿ ಮೀರಿದ ಕಾಮಗಾರಿಗಳ ವಿವರ / Combined Overdue Works by {view_choice}"
)

officer_options = summary_df[view_choice].tolist()

officer_stats_map = {
    row[view_choice]: (
        int(row["overdue_works"]),
        int(row["total_works"]),
        float(row["overdue_pct"]),
    )
    for _, row in summary_df.iterrows()
}

selected_officer = st.selectbox(
    f"Select {view_choice} to view and download all their Overdue Works across all worksheets:",
    options=officer_options,
    index=0 if officer_options else None,
    format_func=lambda name: (
        f"{name}  —  ({officer_stats_map[name][0]} Overdue / "
        f"{officer_stats_map[name][1]} Total | {officer_stats_map[name][2]:.1f}%)"
    ),
)

if selected_officer:
    if overdue_full_df.empty:
        st.info("No overdue works found across the 23 worksheets.")
    else:
        officer_overdue_df = overdue_full_df[
            overdue_full_df[full_filter_col] == selected_officer
        ].copy()

        internal_cols = [c for c in ["_so_name", "_aee_name"] if c in officer_overdue_df.columns]
        officer_overdue_df.drop(columns=internal_cols, inplace=True)

        extra_unnamed = [
            c
            for c in officer_overdue_df.columns
            if c.startswith("Unnamed_Col_") and officer_overdue_df[c].eq("").all()
        ]
        if extra_unnamed:
            officer_overdue_df.drop(columns=extra_unnamed, inplace=True)

        officer_overdue_df = officer_overdue_df.reset_index(drop=True)
        officer_overdue_df.index = officer_overdue_df.index + 1
        officer_overdue_df.index.name = "Sl. No."

        if officer_overdue_df.empty:
            st.success(f"🎉 **{selected_officer}** has **0** overdue works!")
        else:
            sheets_involved = officer_overdue_df["Worksheet / ಯೋಜನೆ"].unique().tolist()
            st.markdown(
                f"**Selected {view_choice}:** `{selected_officer}` &nbsp;|&nbsp; "
                f"**Total Overdue Works:** `{len(officer_overdue_df)}` &nbsp;|&nbsp; "
                f"**Worksheets:** `{', '.join(sheets_involved)}`"
            )

            safe_name = re.sub(r"[^\w\-]+", "_", selected_officer).strip("_") or "Officer"
            excel_bytes = convert_df_to_excel_bytes(officer_overdue_df.reset_index())
            csv_bytes = officer_overdue_df.reset_index().to_csv(index=False).encode("utf-8-sig")

            dl_col1, dl_col2, _ = st.columns([1.5, 1.5, 5])
            with dl_col1:
                st.download_button(
                    label="📥 Download Excel (.xlsx)",
                    data=excel_bytes,
                    file_name=f"Overdue_Works_{safe_name}_{today_str}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                )
            with dl_col2:
                st.download_button(
                    label="📥 Download CSV (.csv)",
                    data=csv_bytes,
                    file_name=f"Overdue_Works_{safe_name}_{today_str}.csv",
                    mime="text/csv",
                    use_container_width=True,
                )

            st.dataframe(
                officer_overdue_df,
                use_container_width=True,
                height=min(600, max(220, (len(officer_overdue_df) + 1) * 42)),
            )

with st.expander("🔍 Inspect Worksheet Loading Status (23 Target Sheets)"):
    st.dataframe(diag_df, use_container_width=True)
